#!/usr/bin/env python3
"""Join cluster-level summaries/centers onto language-level results.

Summary non-key columns (now including centers) are prefixed with "cluster-".
The summary's overall row is not a cluster and is excluded from language joins.
Legacy separate-center inputs are supported, but integrated summaries take priority.
Stage-prefixed reports take priority; unprefixed reports remain supported.
Write each stage's results to meta_data_pretrain_joined or
meta_data_posttrain_joined. The full outer cross-stage join stays in memory.
Cluster CSV joins use only the standard library. Original CSV files are never modified.

For both stages, also left-join Glottolog metadata and WALS 81A/22A/98A
onto the freshly combined language table, writing four independent CSVs per stage.
Population and country-level regional joins use LinguaMeta and UN M49,
with independent CSVs and an audit alongside the four feature tables.
Reading the linguistic XLSX requires openpyxl; no plotting modules are imported.
Workbook pre-/post-training columns are deliberately ignored.
"""

import argparse
import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile


def read_csv(path):
    """Read values as strings to preserve precision, blanks, and language IDs."""
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, strict=True)
        columns = reader.fieldnames
        if not columns or len(columns) != len(set(columns)):
            raise ValueError(f"{path}: missing or duplicate column names")
        keys = [c for c in columns if c.strip().casefold() == "cluster"]
        if len(keys) != 1:
            raise ValueError(f"{path}: expected exactly one Cluster column")
        rows = list(reader)
        for number, row in enumerate(rows, start=2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{path}: incorrect column count at record {number}")
    return columns, keys[0], rows


def cluster_id(value, path):
    """Normalize integer keys, including CSV spellings such as '1.0'."""
    try:
        number = Decimal(value.strip())
    except (InvalidOperation, AttributeError):
        raise ValueError(f"{path}: invalid Cluster value {value!r}") from None
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"{path}: Cluster must be an integer, got {value!r}")
    return int(number)


def index_clusters(rows, key, path, *, skip_overall=False):
    """Enforce many-to-one keys; optionally omit the descriptive overall row."""
    indexed = {}
    overall_seen = False
    for row in rows:
        if skip_overall and row[key].strip().casefold() in {"整体", "overall"}:
            if overall_seen:
                raise ValueError(f"{path}: duplicate overall summary row")
            overall_seen = True
            continue
        cluster = cluster_id(row[key], path)
        if cluster in indexed:
            raise ValueError(f"{path}: duplicate Cluster key {cluster}")
        indexed[cluster] = row
    return indexed


def prepare_join(folder, stage):
    # Choose one naming scheme for the entire stage; never mix new and stale reports.
    prefix = f"{stage}_"
    if not any((folder / f"{prefix}{name}").exists()
               for name in ("cluster_results.csv", "cluster_summary.csv")):
        prefix = ""
    result_path = folder / f"{prefix}cluster_results.csv"
    summary_path = folder / f"{prefix}cluster_summary.csv"
    result_cols, result_key, results = read_csv(result_path)
    summary_cols, summary_key, summaries = read_csv(summary_path)
    summaries = index_clusters(summaries, summary_key, summary_path, skip_overall=True)

    center_fields = [c.removesuffix(" log10") for c in result_cols if c.endswith(" log10")]
    center_fields += [c for c in result_cols if c.endswith("/proportion")]
    embedded_centers = bool(center_fields) and all(c in summary_cols for c in center_fields)
    summary_extra = [c for c in summary_cols if c != summary_key]
    if embedded_centers:
        # Ignore stale legacy center files, even if still present in the folder.
        center_path = summary_path
        center_extra = []
        centers = {
            cluster: {c: row[c] for c in center_fields}
            for cluster, row in summaries.items()
            if cluster != -1 or any(row[c] != "" for c in center_fields)
        }
    else:
        if any(c in summary_cols for c in center_fields):
            raise ValueError(f"{summary_path}: incomplete integrated center columns")
        center_path = folder / f"{prefix}cluster_centers.csv"
        if not center_path.is_file():
            center_path = folder / f"{prefix}cluster_center.csv"
        center_cols, center_key, center_rows = read_csv(center_path)
        centers = index_clusters(center_rows, center_key, center_path)
        center_extra = [c for c in center_cols if c != center_key]

    # Prefix every non-key cluster field once, preserving the old join naming.
    output_cols = result_cols + [
        f"cluster-{c}" for c in summary_extra + center_extra
    ]
    conflicts = [c for c, count in Counter(output_cols).items() if count > 1]
    if conflicts:
        raise ValueError(f"{folder}: output column name collisions: {conflicts}")

    joined = []
    blank_center_rows = 0
    for row in results:
        cluster = cluster_id(row[result_key], result_path)
        if cluster not in summaries:
            raise ValueError(f"{summary_path}: no summary for Cluster {cluster}")
        if cluster not in centers and cluster != -1:
            raise ValueError(f"{center_path}: no center for Cluster {cluster}")
        summary = summaries[cluster]
        center = centers.get(cluster, {})
        blank_center_rows += cluster not in centers

        # Keep every input row in its original order, including Cluster -1.
        merged = dict(row)
        merged.update({f"cluster-{c}": summary[c] for c in summary_extra})
        merged.update({f"cluster-{c}": center.get(c, "") for c in center_extra})
        joined.append(merged)

    return folder / "language_cluster_result.csv", output_cols, joined, blank_center_rows


def index_languages(columns, rows, path):
    """Validate a unique, nonempty language code for a one-to-one stage join.

    Match codes case-insensitively after trimming surrounding whitespace;
    retain original cell values in the output (pre-training spelling wins).
    """
    keys = [c for c in columns if c.strip().casefold() == "language"]
    if len(keys) != 1:
        raise ValueError(f"{path}: expected exactly one Language column")
    key = keys[0]
    indexed = {}
    for number, row in enumerate(rows, start=2):
        value = row.get(key)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"{path}: empty Language at record {number}")
        code = value.strip().casefold()
        if code in indexed:
            raise ValueError(f"{path}: duplicate Language code {value!r}")
        indexed[code] = row
    return key, indexed


def prepare_cross_stage_join(root, pretrain_plan, posttrain_plan):
    """Full outer join the freshly prepared stage tables, not stale CSV files.

    Keep pre-training row order, then append post-training-only languages
    in their original order. Missing stage values remain blank, never -1.
    """
    pre_path, pre_cols, pre_rows, _ = pretrain_plan
    post_path, post_cols, post_rows, _ = posttrain_plan
    pre_key, pre_index = index_languages(pre_cols, pre_rows, pre_path)
    post_key, post_index = index_languages(post_cols, post_rows, post_path)
    pre_extra = [c for c in pre_cols if c != pre_key]
    post_extra = [c for c in post_cols if c != post_key]
    output_cols = (
        ["Language"]
        + [f"pre-{c}" for c in pre_extra]
        + [f"post-{c}" for c in post_extra]
    )
    conflicts = [c for c, count in Counter(output_cols).items() if count > 1]
    if conflicts:
        raise ValueError(f"{root}: cross-stage column name collisions: {conflicts}")

    # The dictionary insertion order mirrors each stage's language row order.
    ordered_codes = list(pre_index) + [c for c in post_index if c not in pre_index]
    joined = []
    for code in ordered_codes:
        pre = pre_index.get(code, {})
        post = post_index.get(code, {})
        language = pre[pre_key] if code in pre_index else post[post_key]
        row = {"Language": language}
        row.update({f"pre-{c}": pre.get(c, "") for c in pre_extra})
        row.update({f"post-{c}": post.get(c, "") for c in post_extra})
        joined.append(row)

    coverage = {
        "matched_languages": len(pre_index.keys() & post_index.keys()),
        "pretrain_only_languages": len(pre_index.keys() - post_index.keys()),
        "posttrain_only_languages": len(post_index.keys() - pre_index.keys()),
    }
    return root / "language_cluster_result.csv", output_cols, joined, coverage


def read_linguistic_features(path):
    """Read only language features, never the workbook's historical cluster fields."""
    from openpyxl import load_workbook

    tables = {}
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in ("Glottolog", "WALS"):
            records = iter(book[sheet].values)
            headers = next(records, ())
            if (not headers or any(not isinstance(c, str) or not c for c in headers)
                    or len(headers) != len(set(headers))):
                raise ValueError(f"{path}/{sheet}: missing or duplicate headers")
            columns = [c for c in headers if not c.startswith(("pre-", "post-", "cluster-"))]
            if sheet == "WALS":
                # Keep WALS metadata and the three features used by the analysis scripts.
                columns = [c for c in columns if not re.match(r"^\d+[A-Z]\s", c)
                           or c.startswith(("81A ", "22A ", "98A "))]
            positions = [headers.index(c) for c in columns]
            rows = [
                {c: "" if values[i] is None else str(values[i])
                 for c, i in zip(columns, positions)}
                for values in records if any(v is not None for v in values)
            ]
            _, indexed = index_languages(columns, rows, f"{path}/{sheet}")
            if any(not re.fullmatch(r"[a-z]{3}", code) for code in indexed):
                raise ValueError(f"{path}/{sheet}: expected ISO 639-3 language codes")
            required = {"Language", "ISO_scope", "ISO_language_type"}
            if sheet == "Glottolog":
                required |= {"classification_status", "family_glottocode", "family_name", "macroareas"}
            if not required.issubset(columns):
                raise ValueError(f"{path}/{sheet}: missing fields {sorted(required - set(columns))}")
            tables[sheet] = columns, indexed
    finally:
        book.close()

    glotto, wals = tables["Glottolog"][1], tables["WALS"][1]
    if glotto.keys() != wals.keys():
        raise ValueError(f"{path}: Glottolog and WALS language-code sets differ")
    for code, row in glotto.items():
        for field in ("ISO_scope", "ISO_language_type"):
            if row[field] != wals[code][field]:
                raise ValueError(f"{path}: conflicting {field} for {code}")
        if row["macroareas"]:
            areas = json.loads(row["macroareas"])
            if not isinstance(areas, list) or not all(isinstance(a, str) for a in areas):
                raise ValueError(f"{path}: expected a macroareas JSON list for {code}")
    return tables


def parse_wals_value(value, feature, code, path):
    """Preserve all observations; only a single distinct category gets a scalar ID."""
    ids, labels = set(), {}
    for line in value.splitlines():
        match = re.fullmatch(r"(\d+)\s+\u2014\s+(.+?)\s+\[([^\]]+)\]", line.strip())
        if not match:
            raise ValueError(f"{path}: unrecognised WALS {feature} value for {code}: {value!r}")
        category = int(match[1])
        maximum = 6 if feature == "98A" else 7
        if not 1 <= category <= maximum:
            raise ValueError(f"{path}: unknown WALS {feature} category {category} for {code}")
        ids.add(category)
        labels.setdefault(category, set()).add(match[2])
    status = "missing" if not ids else "single" if len(ids) == 1 else "conflict"
    prefix = f"WALS_{feature}_"
    return {
        prefix + "category_id": str(next(iter(ids))) if len(ids) == 1 else "",
        prefix + "category_ids": json.dumps(sorted(ids)),
        prefix + "category_labels": json.dumps(
            {str(category): sorted(labels[category]) for category in sorted(ids)},
            ensure_ascii=False,
        ),
        prefix + "status": status,
    }


def prepare_language_feature_join(base_plan, feature_columns, features, output_path,
                                  *, status_column=None):
    """One-to-one left join: preserve current cluster values, row order and all languages."""
    base_path, base_columns, base_rows, _ = base_plan
    key, _ = index_languages(base_columns, base_rows, base_path)
    columns = base_columns + feature_columns
    conflicts = [c for c, count in Counter(columns).items() if count > 1]
    if conflicts:
        raise ValueError(f"{output_path}: feature column name collisions: {conflicts}")
    joined, matched = [], 0
    for row in base_rows:
        code = row[key].strip().casefold()
        matched += code in features
        extra = features.get(code, {})
        merged = dict(row)
        merged.update({c: extra.get(c, "") for c in feature_columns})
        if status_column and code not in features:
            merged[status_column] = "unmatched"
        joined.append(merged)
    stats = {"source_matched": matched, "source_unmatched": len(joined) - matched}
    if status_column:
        stats.update(Counter(row[status_column] for row in joined))
    return output_path, columns, joined, stats


def prepare_linguistic_joins(base_plan, source_path, output_dir):
    """Make four independent feature tables; defer all plot-specific filtering."""
    tables = read_linguistic_features(source_path)
    glotto_columns, glotto = tables["Glottolog"]
    wals_columns, wals = tables["WALS"]
    glotto_columns = [c for c in glotto_columns if c != "Language"]
    plans = [prepare_language_feature_join(
        base_plan, glotto_columns, glotto,
        output_dir / "genealogical_family_analysis_joined.csv",
    )]
    wals_metadata = [c for c in wals_columns
                     if c != "Language" and not re.match(r"^\d+[A-Z]\s", c)]
    for feature, name in (("81A", "word_order_typology_analysis"),
                          ("22A", "verbal_synthesis"), ("98A", "case_alignment")):
        candidates = [c for c in wals_columns if c.startswith(feature + " ")]
        if len(candidates) != 1:
            raise ValueError(f"{source_path}: expected exactly one WALS {feature} column")
        raw_column = candidates[0]
        parsed_columns = [f"WALS_{feature}_{suffix}"
                          for suffix in ("category_id", "category_ids", "category_labels", "status")]
        features = {}
        for code, row in wals.items():
            # Glottolog macroareas support the original regional contrasts.
            # WALS's own geography stays separately prefixed, without overwriting it.
            merged = {c: glotto[code][c] for c in glotto_columns}
            merged.update({f"wals-{c}": row[c] for c in wals_metadata})
            merged[raw_column] = row[raw_column]
            merged.update(parse_wals_value(row[raw_column], feature, code, source_path))
            features[code] = merged
        columns = (glotto_columns + [f"wals-{c}" for c in wals_metadata]
                   + [raw_column] + parsed_columns)
        plans.append(prepare_language_feature_join(
            base_plan, columns, features, output_dir / f"{name}_joined.csv",
            status_column=f"WALS_{feature}_status",
        ))
    return plans


REGIONS = ("Africa", "Americas", "Asia", "Europe", "Oceania", "Unallocated")
QUANTITIES = {"pre": "Num Unicode Characters", "post": "Num Dialogs"}


def read_source_sheet(book, sheet, key):
    """Skip workbook explanation rows and enforce a unique source-table key."""
    rows = iter(book[sheet].values)
    header = next((row for row in rows if key in row), None)
    if header is None or None in header or len(header) != len(set(header)):
        raise ValueError(f"{sheet}: missing or duplicate table headers")
    records = {}
    for values in rows:
        if all(value is None for value in values):
            continue
        row = dict(zip(header, values))
        code = str(row[key] or "").strip().lower()
        if not code or code in records:
            raise ValueError(f"{sheet}: blank or duplicate {key}: {code!r}")
        records[code] = row
    return list(header), records


def resource_quantity(row, prefix, root):
    """Use reduced quantities, not CSV zeros imputed from unavailable log values."""
    code = row["Language"].strip().lower()
    status = row[f"{prefix}-Cluster Status"]
    path = root / f"meta_data_{'pretrain' if prefix == 'pre' else 'posttrain'}_reduced" / f"{code}.json"
    if status == "Uncovered":
        if path.is_file() or row[f"{prefix}-{QUANTITIES[prefix]} log10"] != "":
            raise ValueError(f"{code}: {prefix} Uncovered conflicts with retained data")
        return 0.0, "Uncovered"
    if (status == "" and row[f"{prefix}-{QUANTITIES[prefix]} log10"] == ""
            and row.get(f"{prefix}-Cluster", "") == ""):
        return None, "No stage record"  # An absent stage row is unknown, not Uncovered.
    if status not in ("Clustered", "Insufficient data"):
        raise ValueError(f"{code}: missing or invalid {prefix} Cluster Status: {status!r}")
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    quantity = data.get(QUANTITIES[prefix])
    if quantity is None:
        if status == "Clustered":
            raise ValueError(f"{code}: fitted {prefix} quantity is unavailable")
        return None, "Unavailable"
    if (isinstance(quantity, bool) or not isinstance(quantity, (int, float))
            or not math.isfinite(quantity) or quantity < 0):
        raise ValueError(f"{path}: invalid quantity {quantity!r}")
    if status == "Clustered":
        saved_log = float(row[f"{prefix}-{QUANTITIES[prefix]} log10"])
        if quantity <= 0 or not math.isclose(saved_log, math.log10(quantity), abs_tol=1e-9):
            raise ValueError(f"{code}: {prefix} saved cluster quantity differs from reduced data")
    return quantity, "Known"


def regional_population(data, population, countries):
    """Deduplicate identical script-level estimates; leave conflicts unallocated."""
    by_country = defaultdict(set)
    for record in data.get("language_script_locale", []):
        country = ((record.get("locale") or {}).get("iso_3166_code") or "").upper()
        amount = (record.get("speaker_data") or {}).get("number_of_speakers")
        if (re.fullmatch(r"[A-Z]{2}", country) and not isinstance(amount, bool)
                and isinstance(amount, (int, float)) and math.isfinite(amount) and amount >= 0):
            by_country[country].add(amount)
    conflicts = {c: sorted(v) for c, v in by_country.items() if len(v) > 1}
    known = {c: next(iter(v)) for c, v in by_country.items() if len(v) == 1}
    if math.fsum(known.values()) > population:
        raise ValueError(f"{data['iso_639_3_code']}: country populations exceed total_population")
    amounts = {
        region: math.fsum(v for c, v in known.items()
                          if countries.get(c.lower(), {}).get("Region Name") == region)
        for region in REGIONS[:-1]
    }
    amounts["Unallocated"] = population - math.fsum(amounts.values())
    weights = {region: amount / population for region, amount in amounts.items()}
    unmapped = sorted(c for c in known
                      if countries.get(c.lower(), {}).get("Region Name") not in REGIONS[:-1])
    return amounts, weights, known, conflicts, unmapped


def prepare_linguameta_joins(base_plan, root, source_path, raw_dir, output_dir):
    """Left-join population/geography to fresh cluster results, without plotting."""
    from openpyxl import load_workbook

    source_path, raw_dir, output_dir = Path(source_path), Path(raw_dir), Path(output_dir)
    if not raw_dir.is_dir():
        raise FileNotFoundError(raw_dir)
    book = load_workbook(source_path, read_only=True, data_only=True)
    try:
        source_columns, languages = read_source_sheet(book, "LinguaMeta", "iso_639_3_code")
        _, countries = read_source_sheet(book, "UN_M49", "ISO-alpha2 Code")
    finally:
        book.close()
    required = {"bcp_47_code", "english_name", "total_population", "locales"}
    if not required.issubset(source_columns):
        raise ValueError(f"{source_path}: missing LinguaMeta columns {required - set(source_columns)}")
    if any(not re.fullmatch(r"[a-z]{3}", code) for code in languages):
        raise ValueError(f"{source_path}: invalid ISO 639-3 code")
    if any("Region Name" not in row for row in countries.values()):
        raise ValueError(f"{source_path}: missing UN_M49 Region Name")

    registry = json.loads((root / "language_639-3_info.json").read_text(encoding="utf-8-sig"))
    _, base_columns, base_rows, _ = base_plan
    population_columns = base_columns + source_columns + [
        "ISO_scope", "ISO_language_type", "ISO_Reference_Name", "Chinese_Name",
        "population_match_status", "population_eligible",
        "pre-Q", "pre-Q Status", "post-Q", "post-Q Status",
    ]
    region_columns = population_columns + [
        "regional_status", "regional_summary_eligible", "raw_language_file",
        "country_populations", "conflicting_country_populations", "unmapped_countries",
        "UN_M49 Countries",
    ] + [f"{region} {field}" for region in REGIONS for field in (
        "population", "weight", "pre-Q", "post-Q",
    )]
    if len(region_columns) != len(set(region_columns)):
        raise ValueError("LinguaMeta join would overwrite an existing column")
    population_rows, region_rows = [], []
    scope_names = {"I": "Individual language", "M": "Macrolanguage", "S": "Special"}
    type_names = {"L": "Living", "E": "Extinct", "A": "Ancient", "H": "Historical",
                  "C": "Constructed", "S": "Special"}
    unavailable = {prefix: [] for prefix in QUANTITIES}
    for base in base_rows:
        code = base["Language"].strip().lower()
        info = registry[code]
        source = languages.get(code)
        row = dict(base)
        row.update({column: source.get(column) if source else None for column in source_columns})
        row.update({
            "iso_639_3_code": code,
            "ISO_scope": scope_names.get(info["ISO_639_3_Scope"]),
            "ISO_language_type": type_names.get(info["ISO_639_3_Type"]),
            "ISO_Reference_Name": info["ISO_Reference_Name"],
            "Chinese_Name": info.get("Chinese_Name"),
            "population_match_status": "Matched" if source else "Unmatched",
        })
        population = row["total_population"]
        if population is not None and (isinstance(population, bool)
                or not isinstance(population, (int, float)) or not math.isfinite(population)
                or population < 0):
            raise ValueError(f"{code}: invalid LinguaMeta population {population!r}")
        row["population_eligible"] = (info["ISO_639_3_Scope"] == "I"
                                      and info["ISO_639_3_Type"] == "L"
                                      and population is not None and population > 0)
        for prefix, quantity_key in QUANTITIES.items():
            quantity, status = resource_quantity(base, prefix, root)
            row[f"{prefix}-Q"], row[f"{prefix}-Q Status"] = quantity, status
            # Preserve true log10(1)=0; remove only logs imputed from missing quantities.
            row[f"{prefix}-{quantity_key} log10"] = (
                math.log10(quantity) if quantity is not None and quantity > 0 else None
            )
            if quantity is None:
                unavailable[prefix].append(code)
        population_rows.append(row)
        regional = dict(row)
        regional.update({column: None for column in region_columns if column not in row})
        regional["regional_summary_eligible"] = False
        if not row["population_eligible"]:
            regional["regional_status"] = (
                "Unmatched population" if source is None else "Not individual living"
                if info["ISO_639_3_Scope"] != "I" or info["ISO_639_3_Type"] != "L"
                else "No positive population"
            )
        else:
            bcp = row["bcp_47_code"]
            if not isinstance(bcp, str) or not re.fullmatch(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*", bcp):
                raise ValueError(f"{code}: invalid bcp_47_code {bcp!r}")
            path = raw_dir / f"{bcp}.json"
            regional["raw_language_file"] = str(path.resolve())
            if not path.is_file():
                regional["regional_status"] = "Missing raw language file"
            else:
                data = json.loads(path.read_text(encoding="utf-8-sig"))
                if (data.get("iso_639_3_code") != code or data.get("bcp_47_code") != bcp
                        or data.get("total_population") != population):
                    raise ValueError(f"{path}: raw identity or population differs from workbook")
                amounts, weights, known, conflicts, unmapped = regional_population(data, population, countries)
                regional.update({
                    "regional_status": "Ready", "country_populations": json.dumps(known, sort_keys=True),
                    "conflicting_country_populations": json.dumps(conflicts, sort_keys=True),
                    "unmapped_countries": json.dumps(unmapped),
                    "UN_M49 Countries": json.dumps([countries[c.lower()] for c in sorted(known)
                                                     if c.lower() in countries], ensure_ascii=False),
                })
                # A common complete-case population keeps both stage comparisons consistent.
                regional["regional_summary_eligible"] = all(row[f"{p}-Q"] is not None for p in QUANTITIES)
                for region in REGIONS:
                    regional[f"{region} population"] = amounts[region]
                    regional[f"{region} weight"] = weights[region]
                    for prefix in QUANTITIES:
                        quantity = row[f"{prefix}-Q"]
                        regional[f"{region} {prefix}-Q"] = (
                            quantity * weights[region] if quantity is not None else None
                        )
        region_rows.append(regional)

    eligible = [row for row in region_rows if row["regional_summary_eligible"]]
    population_total = math.fsum(row["total_population"] for row in eligible)
    resource_totals = {p: math.fsum(row[f"{p}-Q"] for row in eligible) for p in QUANTITIES}
    summary_rows = []
    for region in REGIONS:
        population = math.fsum(row[f"{region} population"] for row in eligible)
        summary = {"Region": region, "population": population,
                   "speaker_share_percent": 100 * population / population_total if population_total > 0 else None,
                   "eligible_languages": len(eligible)}
        for prefix in QUANTITIES:
            amount = math.fsum(row[f"{region} {prefix}-Q"] for row in eligible)
            summary[f"{prefix}-Q"] = amount
            summary[f"{prefix}-resource_share_percent"] = (
                100 * amount / resource_totals[prefix] if resource_totals[prefix] > 0 else None
            )
        summary_rows.append(summary)
    for field in ("speaker_share_percent", "pre-resource_share_percent", "post-resource_share_percent"):
        values = [row[field] for row in summary_rows]
        if all(value is None for value in values):
            continue  # An empty/zero denominator defines no share, not a zero share.
        if not math.isclose(math.fsum(values), 100, abs_tol=1e-9):
            raise ValueError(f"{field}: regional shares do not sum to 100%")
    audit = {
        "inputs": {"source_workbook": str(source_path.resolve()), "raw_directory": str(raw_dir.resolve()),
                   "cluster_results": [str(root / f"meta_data_{s}_clustered" / f"{s}_cluster_results.csv")
                                       for s in ("pretrain", "posttrain")],
                   "reduced_directories": [str(root / f"meta_data_{s}_reduced")
                                           for s in ("pretrain", "posttrain")],
                   "registry": str(root / "language_639-3_info.json")},
        "policies": {
            "language_join": "ISO 639-3 left join; retain all current cross-stage rows",
            "regional_scope": "Individual, living languages with positive LinguaMeta population",
            "country_counts": "Identical counts across scripts counted once; conflicting counts excluded",
            "unallocated": "Missing, conflicting, or unmapped country population remains Unallocated",
            "quantity": "Uncovered=0; unavailable reduced quantity=null (never infer Q=1 from imputed CSV zero)",
            "summary": "Same complete-case languages (known Q in both stages) for all three shares",
            "zero_denominator": "Undefined shares remain null; detailed joins are still saved",
            "population_unit": "People; summed language estimates may count multilingual speakers repeatedly",
            "resource_units": {"pre": "Unicode characters", "post": "Reported sample/dialogue equivalents"},
        },
        "language_rows": len(base_rows), "matched_languages": sum(row["population_match_status"] == "Matched" for row in population_rows),
        "population_eligible_languages": sum(row["population_eligible"] for row in population_rows),
        "population_eligible_total": math.fsum(row["total_population"] for row in population_rows
                                               if row["population_eligible"]),
        "regional_status_counts": dict(Counter(row["regional_status"] for row in region_rows)),
        "summary_eligible_languages": len(eligible), "population_total": population_total,
        "summary_excluded_population": math.fsum(row["total_population"] for row in region_rows
                                                 if row["population_eligible"] and not row["regional_summary_eligible"]),
        "undefined_summary_shares": [field for field, value in summary_rows[0].items()
                                     if field.endswith("share_percent") and value is None],
        "resource_totals": resource_totals, "unavailable_quantities": unavailable,
        "summary_excluded_unknown_quantity": [row["Language"] for row in region_rows
                                               if row["regional_status"] == "Ready" and not row["regional_summary_eligible"]],
        "conflicting_languages": [row["Language"] for row in region_rows
                                  if row["conflicting_country_populations"] not in (None, "{}")],
    }
    plans = [
        (output_dir / "population_resource_scatter_joined.csv", population_columns, population_rows,
         {"source_matched": audit["matched_languages"],
          "source_unmatched": len(base_rows) - audit["matched_languages"]}),
        (output_dir / "regional_resource_shares_joined.csv", region_columns, region_rows,
         {"regional_ready": audit["regional_status_counts"].get("Ready", 0),
          "summary_eligible": len(eligible)}),
        (output_dir / "regional_resource_shares_summary.csv", list(summary_rows[0]), summary_rows,
         {"eligible_languages": len(eligible)}),
    ]
    return plans, audit


def write_csv_atomic(path, columns, rows):
    """Finish a temporary file before replacing the output; write Excel-friendly UTF-8."""
    temporary = None
    try:
        with NamedTemporaryFile(
            mode="w", encoding="utf-8-sig", newline="", dir=path.parent,
            prefix=f".{path.name}.", suffix=".tmp", delete=False,
        ) as stream:
            temporary = Path(stream.name)
            writer = csv.DictWriter(stream, fieldnames=columns, lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
        temporary.replace(path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def write_json_atomic(path, content):
    """Atomically save already-validated JSON; never truncate an existing audit."""
    temporary = None
    try:
        with NamedTemporaryFile(
            mode="w", encoding="utf-8", dir=path.parent,
            prefix=f".{path.name}.", suffix=".tmp", delete=False,
        ) as stream:
            temporary = Path(stream.name)
            stream.write(content)
        temporary.replace(path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def project_stage_join(plan, stage, stage_plan, output_dir):
    """Select one stage's fields/row order without changing shared calculations."""
    path, columns, rows, _ = plan
    prefix, other = ("pre", "post") if stage == "pretrain" else ("post", "pre")
    other_region_quantities = {f"{region} {other}-Q" for region in REGIONS}
    selected = {}
    for column in columns:
        if column.startswith(other + "-") or column in other_region_quantities:
            continue
        name = column.removeprefix(prefix + "-")
        for region in REGIONS:
            if name == f"{region} {prefix}-Q":
                name = f"{region} Q"
                break
        selected[column] = name
    if len(set(selected.values())) != len(selected):
        raise ValueError(f"{path}: stage projection creates duplicate column names")

    if "Language" in columns:
        stage_key, stage_index = index_languages(stage_plan[1], stage_plan[2], stage_plan[0])
        _, shared_index = index_languages(columns, rows, path)
        source_rows = []
        for code, stage_row in stage_index.items():
            row = dict(shared_index[code])
            row["Language"] = stage_row[stage_key]
            source_rows.append(row)
    else:
        source_rows = rows  # Regional summaries are shared complete-case aggregates.
    projected = [{name: row[column] for column, name in selected.items()} for row in source_rows]
    return output_dir / path.name, list(selected.values()), projected, {"output_stage": stage}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root", type=Path, default=Path(__file__).resolve().parent,
        help="Project root; default: directory containing this script",
    )
    parser.add_argument(
        "--data-type", choices=("all", "pretrain", "posttrain"), default="all",
        help="all (default): all joins in both stage directories; one stage: its cluster join only",
    )
    parser.add_argument(
        "--check-only", action="store_true",
        help="Validate and join in memory without writing output CSV files",
    )
    parser.add_argument(
        "--linguistic-input", type=Path,
        help="Feature XLSX; default: <root>/linguistic_analysis/language_cluster_glottolog_wals.xlsx",
    )
    parser.add_argument(
        "--output-root", "--linguistic-output-dir", dest="output_root", type=Path,
        help="Parent of meta_data_pretrain_joined/meta_data_posttrain_joined; default: project root",
    )
    parser.add_argument(
        "--skip-linguistic-joins", action="store_true",
        help="Only generate the original cluster joins; linguistic joins require --data-type all",
    )
    parser.add_argument(
        "--linguameta-input", type=Path,
        help="Population/M49 XLSX; default: <root>/linguistic_analysis/LinguaMeta_UN_M49_source_tables.xlsx",
    )
    parser.add_argument(
        "--linguameta-raw-dir", type=Path,
        help="Country speaker JSONs; default: <root>/linguistic_analysis/linguameta/raw/data",
    )
    args = parser.parse_args()
    root = args.root.expanduser().resolve()
    stages = ("pretrain", "posttrain") if args.data_type == "all" else (args.data_type,)

    # Validate all cluster and linguistic joins before writing any output.
    try:
        plans = [prepare_join(root / f"meta_data_{stage}_clustered", stage) for stage in stages]
        combined = (
            prepare_cross_stage_join(root, plans[0], plans[1])
            if args.data_type == "all" else None
        )
        linguistic_input = (args.linguistic_input or root / "linguistic_analysis"
                            / "language_cluster_glottolog_wals.xlsx").expanduser().resolve()
        output_root = (args.output_root or root).expanduser().resolve()
        linguistic_plans = (
            prepare_linguistic_joins(combined, linguistic_input, output_root)
            if combined is not None and not args.skip_linguistic_joins else []
        )
        linguameta_plans, linguameta_audit = [], None
        if combined is not None and not args.skip_linguistic_joins:
            source_path = (args.linguameta_input or root / "linguistic_analysis"
                           / "LinguaMeta_UN_M49_source_tables.xlsx").expanduser().resolve()
            raw_dir = (args.linguameta_raw_dir or root / "linguistic_analysis"
                       / "linguameta" / "raw" / "data").expanduser().resolve()
            linguameta_plans, linguameta_audit = prepare_linguameta_joins(
                combined, root, source_path, raw_dir, output_root,
            )
        output_plans, audit_plans = [], []
        for stage, stage_plan in zip(stages, plans):
            output_dir = output_root / f"meta_data_{stage}_joined"
            output_plans.append((output_dir / stage_plan[0].name, stage_plan[1], stage_plan[2],
                                 {"rows_without_center": stage_plan[3]}))
            output_plans.extend(project_stage_join(plan, stage, stage_plan, output_dir)
                                for plan in linguistic_plans + linguameta_plans)
            if linguameta_audit is not None:
                # Retain both-stage provenance/exclusions for the unchanged common sample.
                audit = dict(linguameta_audit, output_stage=stage,
                             output_language_rows=len(stage_plan[2]))
                content = json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
                audit_plans.append((output_dir / "linguameta_join_summary.json", content))
            ancestor = output_dir
            while not ancestor.exists() and ancestor.parent != ancestor:
                ancestor = ancestor.parent
            if not ancestor.is_dir():
                raise ValueError(f"{output_dir}: output directory has a non-directory ancestor")
        output_paths = [plan[0] for plan in output_plans + audit_plans]
        if any(path.exists() and not path.is_file() for path in output_paths):
            raise ValueError("An output target already exists but is not a file")
        for path, columns, rows, stats in output_plans:
            if not args.check_only:
                path.parent.mkdir(parents=True, exist_ok=True)
                write_csv_atomic(path, columns, rows)
            action = "CHECKED" if args.check_only else "SAVED"
            details = ", ".join(f"{key}={value}" for key, value in stats.items())
            print(f"{action} {path} | rows={len(rows)}, columns={len(columns)}, {details}")
        for audit_path, content in audit_plans:
            if not args.check_only:
                write_json_atomic(audit_path, content)
            action = "CHECKED" if args.check_only else "SAVED"
            print(f"{action} {audit_path}")
    except (OSError, ValueError, csv.Error, KeyError, ImportError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
